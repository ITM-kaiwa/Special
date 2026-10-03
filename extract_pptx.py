"""
PowerPoint 内容抽出スクリプト
Z:\Khối Đào tạo\3. DAO TAO DAC BIET - 特別教育 内の全 .pptx/.pptm を読み取り、
スライドごとのテキストを JSON で出力する。
"""
import os, json, sys, re
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = r"Z:\Khối Đào tạo\3. DAO TAO DAC BIET - 特別教育"
OUT = r"C:\Users\Admin\Projects\Special\extracted"

os.makedirs(OUT, exist_ok=True)

def extract_text_from_shape(shape):
    """シェイプからテキストを抽出"""
    texts = []
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            line = ""
            for run in para.runs:
                line += run.text
            if line.strip():
                texts.append(line.strip())
    return texts

def extract_table(shape):
    """テーブルからデータを抽出"""
    if not shape.has_table:
        return None
    table = shape.table
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cells.append(cell.text.strip())
        rows.append(cells)
    return rows

def extract_pptx(filepath):
    """PPTXファイルからスライドごとのデータを抽出"""
    try:
        prs = Presentation(filepath)
    except Exception as e:
        print(f"  ERROR opening: {e}", file=sys.stderr)
        return None

    slides_data = []
    for i, slide in enumerate(prs.slides):
        slide_info = {
            "slide_number": i + 1,
            "texts": [],
            "tables": [],
            "notes": ""
        }
        
        for shape in slide.shapes:
            # テキスト
            texts = extract_text_from_shape(shape)
            if texts:
                slide_info["texts"].extend(texts)
            
            # テーブル
            if shape.has_table:
                table_data = extract_table(shape)
                if table_data:
                    slide_info["tables"].append(table_data)
            
            # グループ内のシェイプ
            if shape.shape_type == 6:  # GROUP
                try:
                    for child in shape.shapes:
                        texts = extract_text_from_shape(child)
                        if texts:
                            slide_info["texts"].extend(texts)
                        if hasattr(child, 'has_table') and child.has_table:
                            table_data = extract_table(child)
                            if table_data:
                                slide_info["tables"].append(table_data)
                except:
                    pass
        
        # ノート
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame:
                slide_info["notes"] = notes_frame.text.strip()
        
        slides_data.append(slide_info)
    
    return {
        "filename": os.path.basename(filepath),
        "total_slides": len(slides_data),
        "slides": slides_data
    }

def main():
    files = sorted([
        f for f in os.listdir(SRC)
        if f.endswith(('.pptx', '.pptm')) and not f.startswith('~')
    ])
    
    print(f"Found {len(files)} PowerPoint files")
    
    summary = []
    
    for fname in files:
        filepath = os.path.join(SRC, fname)
        print(f"Processing: {fname} ...", end=" ", flush=True)
        
        data = extract_pptx(filepath)
        if data is None:
            print("SKIPPED")
            continue
        
        # 安全なファイル名
        safe_name = re.sub(r'[^\w\-]', '_', os.path.splitext(fname)[0])
        out_path = os.path.join(OUT, f"{safe_name}.json")
        
        with open(out_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        print(f"OK ({data['total_slides']} slides)")
        summary.append({
            "filename": fname,
            "safe_name": safe_name,
            "slides": data["total_slides"],
            "output": out_path
        })
    
    # サマリー出力
    summary_path = os.path.join(OUT, "_summary.json")
    with open(summary_path, 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    
    print(f"\nDone! {len(summary)} files processed.")
    print(f"Summary: {summary_path}")

if __name__ == "__main__":
    main()
